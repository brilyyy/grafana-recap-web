from __future__ import annotations

import io
from pathlib import Path
from typing import Any, List

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from ._finder import ShapeFinder
from ._models import ImageReplacement, Replacement, Result, TableData, TableStyle


class TextHandler:
    """Handle text replacement in shapes"""

    TYPE = "TEXT"

    def can_handle(self, shape: Any) -> bool:
        return hasattr(shape, "text_frame")

    def replace(self, shape: Any, data: Replacement) -> Result:
        """Replace placeholder text in shape, preserving first-run formatting."""
        try:
            if not self.can_handle(shape):
                return Result(
                    False, self.TYPE, data.placeholder, "Shape has no text frame"
                )

            text_frame = shape.text_frame
            if data.placeholder not in text_frame.text:
                return Result(
                    False, self.TYPE, data.placeholder, "Placeholder not found"
                )

            replaced = False
            for paragraph in text_frame.paragraphs:
                if data.placeholder not in paragraph.text:
                    continue

                new_text = paragraph.text.replace(data.placeholder, data.value)
                runs = paragraph.runs

                if runs:
                    # Keep first run's character formatting; remove excess runs from XML.
                    for run in runs[1:]:
                        run._r.getparent().remove(run._r)
                    runs[0].text = new_text
                    if data.style:
                        data.style.apply(runs[0], paragraph)
                else:
                    run = paragraph.add_run()
                    run.text = new_text
                    if data.style:
                        data.style.apply(run, paragraph)

                replaced = True

            if replaced:
                short = data.value[:50]
                suffix = "..." if len(data.value) > 50 else ""
                return Result(
                    True,
                    self.TYPE,
                    data.placeholder,
                    f"Replaced with '{short}{suffix}'",
                )
            return Result(False, self.TYPE, data.placeholder, "Replacement failed")

        except Exception as e:
            return Result(False, self.TYPE, data.placeholder, str(e))

    def replace_global(self, pres: Any, data: Replacement) -> List[Result]:
        """Replace placeholder across all slides"""
        results: List[Result] = []
        for _slide_num, shape in ShapeFinder.all_shapes(pres):
            if self.can_handle(shape):
                r = self.replace(shape, data)
                if r.success:
                    results.append(r)
        return results


class ImageHandler:
    """Handle image replacement"""

    TYPE = "IMAGE"

    def can_handle(self, shape: Any) -> bool:
        try:
            return shape._element.blipFill is not None
        except AttributeError:
            return False

    def replace(self, shape: Any, data: ImageReplacement) -> Result:
        """Replace image using the slide part's public get_or_add_image_part API."""
        try:
            if not self.can_handle(shape):
                return Result(
                    False, self.TYPE, data.placeholder, "Shape is not a picture"
                )

            left, top, width, height = shape.left, shape.top, shape.width, shape.height

            # get_or_add_image_part accepts a path or a file-like stream.
            slide_part = shape.part
            if isinstance(data.image_path, bytes):
                image_source: Any = io.BytesIO(data.image_path)
                label = "<bytes>"
            else:
                p = Path(data.image_path)  # type: ignore
                if not p.exists():
                    return Result(
                        False, self.TYPE, data.placeholder, f"Image not found: {p}"
                    )
                image_source = str(p)
                label = p.name

            _, rId = slide_part.get_or_add_image_part(image_source)
            shape._element.blipFill.blip.set(qn("r:embed"), rId)

            shape.left, shape.top = left, top
            if not data.preserve_aspect:
                shape.width, shape.height = width, height

            return Result(True, self.TYPE, data.placeholder, f"Replaced with {label}")

        except Exception as e:
            return Result(False, self.TYPE, data.placeholder, str(e))


class TableHandler:
    """Handle table data replacement"""

    TYPE = "TABLE"

    def can_handle(self, shape: Any) -> bool:
        return shape.shape_type == MSO_SHAPE_TYPE.TABLE

    def replace(
        self, shape: Any, data: TableData, style: TableStyle | None = None
    ) -> Result:
        """Replace table contents. Template table must be pre-sized to match data dimensions."""
        try:
            if not self.can_handle(shape):
                return Result(False, self.TYPE, "TABLE", "Shape is not a table")

            table = shape.table
            current_rows = len(table.rows)
            current_cols = len(table.columns)
            needed_rows = len(data.rows) + 1  # +1 for header
            needed_cols = len(data.headers)

            if current_rows != needed_rows or current_cols != needed_cols:
                return Result(
                    False,
                    self.TYPE,
                    "TABLE",
                    f"Table size mismatch: {current_rows}x{current_cols} "
                    f"vs needed {needed_rows}x{needed_cols}. "
                    f"Pre-size your template table.",
                )

            header_style = style.header_style if style else None
            row_style = style.row_style if style else None

            for col_idx, header in enumerate(data.headers):
                self._set_cell_text(table.cell(0, col_idx), header, header_style)

            for row_idx, row_data in enumerate(data.rows, 1):
                for col_idx, cell_text in enumerate(row_data):
                    self._set_cell_text(
                        table.cell(row_idx, col_idx), str(cell_text), row_style
                    )

            if style is not None and (
                style.word_wrap or style.cell_margin_pt is not None
            ):
                self._apply_cell_fit(table, style)

            return Result(
                True, self.TYPE, "TABLE", f"Updated {needed_rows}x{needed_cols} table"
            )

        except Exception as e:
            return Result(False, self.TYPE, "TABLE", str(e))

    @staticmethod
    def _set_cell_text(cell: Any, value: str, style: Any = None) -> None:
        """Write text into a cell, preserving existing run formatting unless style overrides it.

        python-pptx's cell.text setter is destructive — it replaces the entire txBody
        with a bare run that has no explicit character properties, causing the text to
        inherit a larger theme default size. This method writes only into the <a:t>
        element of the first existing run, keeping all <a:rPr> attributes intact.
        If a TextStyle is provided it is applied on top of the preserved formatting.
        """
        tf = cell.text_frame
        paragraph = tf.paragraphs[0]
        runs = paragraph.runs
        if runs:
            # Overwrite text of the first run; drop any excess runs.
            runs[0].text = value
            for run in runs[1:]:
                run._r.getparent().remove(run._r)
            if style is not None:
                style.apply(runs[0], paragraph)
        else:
            # No runs — create one (formatting will inherit from cell/table style).
            run = paragraph.add_run()
            run.text = value
            if style is not None:
                style.apply(run, paragraph)

    @staticmethod
    def _apply_cell_fit(table: Any, style: TableStyle) -> None:
        for row in table.rows:
            for cell in row.cells:
                tf = cell.text_frame
                if style.word_wrap:
                    tf.word_wrap = True
                    for p in tf.paragraphs:
                        p.space_before = Pt(0)
                        p.space_after = Pt(0)
                if style.cell_margin_pt is not None:
                    m = Pt(style.cell_margin_pt)
                    tf.margin_left = m
                    tf.margin_right = m
                    tf.margin_top = m
                    tf.margin_bottom = m
