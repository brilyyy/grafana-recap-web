from __future__ import annotations

import builtins
import io
import warnings
from copy import deepcopy
from datetime import datetime
from typing import Any, Dict, List, Literal, Optional

import pptx.presentation as _pptx_pres
from pptx import Presentation as _make_pres
from pptx.util import Inches

from ._engine import TemplateEngine
from ._exceptions import BPTXError, TemplateError
from ._models import BatchResult, Result, TableData, TableStyle, TextStyle
from ._types import ImageSource, PathLike


class BPTX:
    """
    Business PowerPoint Template Framework.
    Fluent API for template-based PowerPoint automation.
    """

    def __init__(self) -> None:
        self._pres: Optional[_pptx_pres.Presentation] = None
        self._engine: Optional[TemplateEngine] = None
        self._log: List[Result] = []
        self._template_path: Optional[Any] = None  # pathlib.Path

    # -------------------------------------------------------------------------
    # LIFECYCLE
    # -------------------------------------------------------------------------

    def load(self, path: PathLike) -> BPTX:
        """Load an existing .pptx template"""
        from pathlib import Path

        self._template_path = Path(path)
        if not self._template_path.exists():
            raise TemplateError(f"Template not found: {path}")
        self._pres = _make_pres(str(self._template_path))
        self._engine = TemplateEngine(self._pres)
        return self

    def new(self, width: float = 10, height: float = 7.5) -> BPTX:
        """Create a blank presentation"""
        self._pres = _make_pres()
        self._pres.slide_width = Inches(width)
        self._pres.slide_height = Inches(height)
        self._engine = TemplateEngine(self._pres)
        return self

    def save(self, path: Optional[PathLike] = None) -> BPTX:
        """Save presentation to disk"""
        from pathlib import Path

        if not self._pres:
            raise BPTXError("No presentation loaded")

        save_path = Path(path) if path else self._template_path
        if not save_path:
            raise BPTXError("No save path specified")

        if save_path == self._template_path and path is None:
            warnings.warn(
                "Saving to the original template path. Pass an explicit path to avoid overwriting the template.",
                stacklevel=2,
            )

        self._pres.save(str(save_path))
        return self

    def bytes(self) -> builtins.bytes:
        """Return presentation as raw bytes"""
        if not self._pres:
            raise BPTXError("No presentation loaded")
        stream = io.BytesIO()
        self._pres.save(stream)
        return stream.getvalue()

    def close(self) -> None:
        """Release resources"""
        self._pres = None
        self._engine = None

    def __enter__(self) -> BPTX:
        return self

    def __exit__(self, *args: Any) -> Literal[False]:
        self.close()
        return False

    # -------------------------------------------------------------------------
    # TEXT REPLACEMENT
    # -------------------------------------------------------------------------

    def replace_text(
        self, placeholder: str, value: Any, style: Optional[TextStyle] = None
    ) -> BPTX:
        """Replace {{placeholder}} with value across all slides"""
        if not self._engine:
            raise BPTXError("No template loaded")

        str_value = (
            value.strftime("%Y-%m-%d") if isinstance(value, datetime) else str(value)
        )
        results = self._engine.replace(placeholder, str_value, style)
        self._log.extend(results)
        return self

    def replace_text_many(self, mapping: Dict[str, Any]) -> BatchResult:
        """
        Replace multiple placeholders from a dict across all slides.
        Returns a BatchResult summarising all operations.
        """
        if not self._engine:
            raise BPTXError("No template loaded")

        all_results: List[Result] = []
        for placeholder, value in mapping.items():
            str_value = (
                value.strftime("%Y-%m-%d")
                if isinstance(value, datetime)
                else str(value)
            )
            results = self._engine.replace(placeholder, str_value)
            all_results.extend(results)

        self._log.extend(all_results)
        succeeded = sum(1 for r in all_results if r.success)
        return BatchResult(
            total=len(all_results),
            succeeded=succeeded,
            failed=len(all_results) - succeeded,
            results=all_results,
        )

    def replace_text_on_slide(
        self, slide_num: int, placeholder: str, value: str
    ) -> BPTX:
        """Replace placeholder on a specific slide only"""
        if not self._engine:
            raise BPTXError("No template loaded")
        results = self._engine.replace_on_slide(slide_num, placeholder, value)
        self._log.extend(results)
        return self

    def replace_many_on_slide(
        self, slide_num: int, mapping: Dict[str, Any]
    ) -> BatchResult:
        """
        Replace multiple placeholders on a specific slide.
        Returns a BatchResult summarising all operations.
        """
        if not self._engine:
            raise BPTXError("No template loaded")

        all_results: List[Result] = []
        for placeholder, value in mapping.items():
            str_value = (
                value.strftime("%Y-%m-%d")
                if isinstance(value, datetime)
                else str(value)
            )
            results = self._engine.replace_on_slide(slide_num, placeholder, str_value)
            all_results.extend(results)

        self._log.extend(all_results)
        succeeded = sum(1 for r in all_results if r.success)
        return BatchResult(
            total=len(all_results),
            succeeded=succeeded,
            failed=len(all_results) - succeeded,
            results=all_results,
        )

    # -------------------------------------------------------------------------
    # IMAGE
    # -------------------------------------------------------------------------

    def replace_image(self, shape_name: str, image_path: ImageSource) -> BPTX:
        """Replace image in named shape. Accepts a file path or raw PNG bytes."""
        if not self._engine:
            raise BPTXError("No template loaded")
        r = self._engine.replace_image(shape_name, image_path)
        self._log.append(r)
        return self

    def replace_image_on_slide(
        self, slide_num: int, shape_name: str, image_path: ImageSource
    ) -> BPTX:
        if not self._engine:
            raise BPTXError("No template loaded")
        r = self._engine.replace_image_on_slide(slide_num, shape_name, image_path)
        self._log.append(r)
        return self

    # -------------------------------------------------------------------------
    # TABLE
    # -------------------------------------------------------------------------

    def replace_table(
        self, shape_name: str, data: TableData, style: Optional[TableStyle] = None
    ) -> BPTX:
        """Replace table data in named shape. Pass a TableStyle to override header/row formatting."""
        if not self._engine:
            raise BPTXError("No template loaded")
        r = self._engine.replace_table(shape_name, data, style)
        self._log.append(r)
        return self

    def replace_table_on_slide(
        self,
        slide_num: int,
        shape_name: str,
        data: TableData,
        style: Optional[TableStyle] = None,
    ) -> BPTX:
        """Replace table data in named shape on a specific slide."""
        if not self._engine:
            raise BPTXError("No template loaded")
        r = self._engine.replace_table_on_slide(slide_num, shape_name, data, style)
        self._log.append(r)
        return self

    def replace_table_rows(self, shape_name: str, data: List[List[str]]) -> BPTX:
        """Replace table from a 2D array (first row = headers)"""
        if not data or not data[0]:
            raise BPTXError("Table data cannot be empty")
        return self.replace_table(shape_name, TableData(headers=data[0], rows=data[1:]))

    # -------------------------------------------------------------------------
    # SLIDE MANIPULATION
    # -------------------------------------------------------------------------

    def copy_slide(self, slide_num: int) -> int:
        """Duplicate a slide, appending it at the end. Returns the new slide number."""
        if not self._pres:
            raise BPTXError("No presentation loaded")
        return self._copy_slide_impl(slide_num)

    def copy_slide_to(self, slide_from: int, slide_to: int) -> int:
        """
        Duplicate a slide and insert the copy at a specific position.
        Returns the new slide number (equal to slide_to).
        """
        if not self._pres:
            raise BPTXError("No presentation loaded")

        total = len(self._pres.slides)
        if not (1 <= slide_from <= total):
            raise BPTXError(f"slide_from {slide_from} out of range (1–{total})")
        if not (1 <= slide_to <= total + 1):
            raise BPTXError(f"slide_to {slide_to} out of range (1–{total + 1})")

        self._copy_slide_impl(slide_from)

        # Move the newly appended slide (now at the end) to slide_to position.
        sldIdLst = self._pres.slides._sldIdLst
        last = sldIdLst[-1]
        sldIdLst.remove(last)
        sldIdLst.insert(slide_to - 1, last)

        return slide_to

    def _copy_slide_impl(self, slide_num: int) -> int:
        """Copy a slide to the end. Returns the new slide number."""
        assert self._pres is not None

        source_slide = self._pres.slides[slide_num - 1]
        new_slide = self._pres.slides.add_slide(source_slide.slide_layout)

        source_spTree = source_slide.shapes._spTree
        new_spTree = new_slide.shapes._spTree

        # Replace new slide's shape tree wholesale to avoid duplicate layout shapes.
        for child in list(new_spTree):
            new_spTree.remove(child)
        for child in source_spTree:
            new_spTree.append(deepcopy(child))

        # Propagate notes slide if present.
        if source_slide.has_notes_slide:
            source_notes_xml = deepcopy(source_slide.notes_slide._element)
            target_notes_el = new_slide.notes_slide._element
            parent = target_notes_el.getparent()
            if parent is not None:
                parent.replace(target_notes_el, source_notes_xml)

        return len(self._pres.slides)

    def remove_slide(self, slide_num: int) -> BPTX:
        """Remove a slide by number (1-indexed)"""
        if not self._pres:
            raise BPTXError("No presentation loaded")

        # python-pptx has no public remove_slide API; this is the canonical workaround.
        rId = self._pres.slides._sldIdLst[slide_num - 1].rId
        self._pres.part.drop_rel(rId)
        del self._pres.slides._sldIdLst[slide_num - 1]
        return self

    # -------------------------------------------------------------------------
    # INTROSPECTION
    # -------------------------------------------------------------------------

    def find_placeholders(self) -> Dict[int, List[str]]:
        """Return {slide_num: [placeholder, ...]} for all {{placeholders}} found"""
        if not self._engine:
            raise BPTXError("No template loaded")
        return self._engine.scan_placeholders()

    def log(self) -> List[Result]:
        """Return a copy of the operation log"""
        return self._log.copy()

    def summary(self) -> str:
        """Human-readable summary of the current presentation state"""
        if not self._pres:
            return "No presentation loaded"

        lines = [
            f"Slides: {len(self._pres.slides)}",
            f"Operations: {len(self._log)}",
        ]
        if self._log:
            rate = sum(1 for r in self._log if r.success) / len(self._log) * 100
            lines.append(f"Success rate: {rate:.1f}%")
        else:
            lines.append("No operations")
        return "\n".join(lines)
