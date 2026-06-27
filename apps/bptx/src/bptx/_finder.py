from __future__ import annotations

from typing import Any, Iterator, Optional, Tuple

import pptx.presentation as _pptx_pres


class ShapeFinder:
    """Find shapes by various criteria"""

    @staticmethod
    def by_name(slide: Any, name: str) -> Optional[Any]:
        """Find shape by name"""
        for shape in slide.shapes:
            if shape.name == name:
                return shape
        return None

    @staticmethod
    def by_placeholder(slide: Any, placeholder: str) -> Iterator[Any]:
        """Find shapes containing placeholder text"""
        for shape in slide.shapes:
            if ShapeFinder._has_text(shape, placeholder):
                yield shape

    @staticmethod
    def _has_text(shape: Any, text: str) -> bool:
        """Check if shape contains text"""
        try:
            if hasattr(shape, "text"):
                return text in shape.text
            if hasattr(shape, "text_frame"):
                return text in shape.text_frame.text
            return False
        except Exception:
            return False

    @staticmethod
    def all_shapes(pres: _pptx_pres.Presentation) -> Iterator[Tuple[int, Any]]:
        """Iterate all shapes across all slides"""
        for slide_num, slide in enumerate(pres.slides, 1):
            for shape in slide.shapes:
                yield (slide_num, shape)
