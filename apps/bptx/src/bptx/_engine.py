from __future__ import annotations

import re
from typing import Any, Dict, List, Optional

import pptx.presentation as _pptx_pres

from ._finder import ShapeFinder
from ._handlers import ImageHandler, TableHandler, TextHandler
from ._models import (
    ImageReplacement,
    Replacement,
    Result,
    TableData,
    TableStyle,
    TextStyle,
)
from ._types import ImageSource


class TemplateEngine:
    """Process template placeholders using registered handlers"""

    PLACEHOLDER_PATTERN = re.compile(r"\{\{(\w+)\}\}")

    def __init__(
        self,
        pres: _pptx_pres.Presentation,
        text_handler: Optional[TextHandler] = None,
        image_handler: Optional[ImageHandler] = None,
        table_handler: Optional[TableHandler] = None,
    ) -> None:
        self._pres = pres
        self._text_handler = text_handler or TextHandler()
        self._image_handler = image_handler or ImageHandler()
        self._table_handler = table_handler or TableHandler()

    def scan_placeholders(self) -> Dict[int, List[str]]:
        """Find all {{placeholders}} across all slides"""
        placeholders: Dict[int, List[str]] = {}

        for slide_num, slide in enumerate(self._pres.slides, 1):
            found: List[str] = []
            for shape in slide.shapes:
                text = self._get_shape_text(shape)
                if text:
                    found.extend(self.PLACEHOLDER_PATTERN.findall(text))
            if found:
                placeholders[slide_num] = list(set(found))

        return placeholders

    def _get_shape_text(self, shape: Any) -> Optional[str]:
        try:
            if hasattr(shape, "text"):
                return shape.text
            if hasattr(shape, "text_frame"):
                return shape.text_frame.text
            return None
        except Exception:
            return None

    def replace(
        self, placeholder: str, value: str, style: Optional[TextStyle] = None
    ) -> List[Result]:
        """Replace placeholder across all slides"""
        replacement = Replacement(placeholder, value, style)
        return self._text_handler.replace_global(self._pres, replacement)

    def replace_on_slide(
        self, slide_num: int, placeholder: str, value: str
    ) -> List[Result]:
        """Replace placeholder on a specific slide, all matching shapes"""
        if slide_num < 1 or slide_num > len(self._pres.slides):
            return [
                Result(
                    False, "TEMPLATE", placeholder, f"Invalid slide number: {slide_num}"
                )
            ]

        slide = self._pres.slides[slide_num - 1]
        replacement = Replacement(placeholder, value)
        results: List[Result] = []

        for shape in slide.shapes:
            if self._text_handler.can_handle(shape):
                r = self._text_handler.replace(shape, replacement)
                if r.success:
                    results.append(r)

        if not results:
            return [
                Result(False, "TEMPLATE", placeholder, "Placeholder not found on slide")
            ]
        return results

    def replace_image(self, shape_name: str, image_path: ImageSource) -> Result:
        """Replace image in named shape"""
        for _slide_num, shape in ShapeFinder.all_shapes(self._pres):
            if shape.name == shape_name:
                data = ImageReplacement(shape_name, image_path)
                return self._image_handler.replace(shape, data)
        return Result(False, "IMAGE", shape_name, "Shape not found")

    def replace_table(
        self, table_shape: str, data: TableData, style: Optional[TableStyle] = None
    ) -> Result:
        """Replace table data in named shape"""
        for _slide_num, shape in ShapeFinder.all_shapes(self._pres):
            if shape.name == table_shape:
                return self._table_handler.replace(shape, data, style)
        return Result(False, "TABLE", table_shape, "Table not found")

    def replace_table_on_slide(
        self,
        slide_num: int,
        shape_name: str,
        data: TableData,
        style: Optional[TableStyle] = None,
    ) -> Result:
        """Replace table data in named shape on a specific slide"""
        if slide_num < 1 or slide_num > len(self._pres.slides):
            return Result(
                False, "TABLE", shape_name, f"Invalid slide number: {slide_num}"
            )
        slide = self._pres.slides[slide_num - 1]
        for shape in slide.shapes:
            if shape.name == shape_name:
                return self._table_handler.replace(shape, data, style)
        return Result(False, "TABLE", shape_name, "Table not found on slide")

    def replace_image_on_slide(
        self, slide_num: int, shape_name: str, image_path: ImageSource
    ) -> Result:
        if slide_num < 1 or slide_num > len(self._pres.slides):
            return Result(
                False, "Image", shape_name, f"Invalid slide number: {slide_num}"
            )
        slide = self._pres.slides[slide_num - 1]
        for shape in slide.shapes:
            if shape.name == shape_name:
                data = ImageReplacement(shape_name, image_path)
                return self._image_handler.replace(shape, data)
        return Result(False, "IMAGE", shape_name, "Image not found on slide")
